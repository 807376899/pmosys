"""Small openpyxl-only helpers for exercising real Excel HTTP payloads."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook


class _Column(list):
    def tolist(self) -> list[Any]:
        return list(self)


class _Loc:
    def __init__(self, frame: "DataFrame") -> None:
        self.frame = frame

    def __getitem__(self, key: tuple[int, str]) -> Any:
        return self.frame.rows[key[0]].get(key[1])


class DataFrame:
    def __init__(self, rows: list[dict[str, Any]] | list[list[Any]], columns: list[str] | None = None) -> None:
        if columns is None:
            columns = list(dict.fromkeys(key for row in rows for key in row)) if rows else []  # type: ignore[union-attr]
        self.columns = columns
        self.rows = [dict(zip(columns, row)) if not isinstance(row, dict) else {column: row.get(column) for column in columns} for row in rows]
        self.loc = _Loc(self)

    def __getitem__(self, column: str) -> _Column:
        return _Column(row.get(column) for row in self.rows)

    def to_excel(self, output: io.BytesIO, index: bool = False, sheet_name: str = "Sheet1") -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(self.columns)
        for row in self.rows:
            sheet.append([row.get(column) for column in self.columns])
        workbook.save(output)


def read_excel(content: io.BytesIO, sheet_name: str | None = None) -> DataFrame:
    workbook = load_workbook(content, data_only=True, read_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return DataFrame([])
    headers = [str(value or "") for value in values[0]]
    return DataFrame([dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)], headers)
